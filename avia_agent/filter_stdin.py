import io
import sys


def _write_utf8(text: str) -> None:
    sys.stdout.buffer.write(text.encode("utf-8", errors="ignore"))


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        txt = page.extract_text() or ""
        parts.append(f"\n@@@PAGE {i}@@@\n{txt}")
    return "\n".join(parts)


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _pptx_to_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    chunks: list[str] = []
    for slide_i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                chunks.append(f"[slide {slide_i}] {text}")
    return "\n".join(chunks)


def _xlsx_to_text(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"\n--- sheet {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip() != ""]
            if vals:
                chunks.append("\t".join(vals))
    return "\n".join(chunks)


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        print("usage: python -m avia_agent.filter_stdin <ext>", file=sys.stderr)
        return 2

    ext = argv[0].lower().lstrip(".")
    data = sys.stdin.buffer.read()

    try:
        if ext == "pdf":
            _write_utf8(_pdf_to_text(data))
            return 0
        if ext == "docx":
            _write_utf8(_docx_to_text(data))
            return 0
        if ext == "pptx":
            _write_utf8(_pptx_to_text(data))
            return 0
        if ext in {"xlsx", "xlsm", "xltx", "xltm"}:
            _write_utf8(_xlsx_to_text(data))
            return 0

        _write_utf8(data.decode("utf-8", errors="ignore"))
        return 0
    except Exception as e:
        print(f"error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
