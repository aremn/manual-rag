import json
import os
import re
import shutil
import subprocess
import sys
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any, Literal

from pydantic import BaseModel, Field


def _safe_abspath(path: str) -> str:
    return os.path.abspath(os.path.expandvars(os.path.expanduser(path)))


def _is_within(root: str, path: str) -> bool:
    root = os.path.normcase(_safe_abspath(root))
    path = os.path.normcase(_safe_abspath(path))
    try:
        common = os.path.commonpath([root, path])
    except ValueError:
        return False
    return common == root


def _truncate(s: str, max_chars: int) -> str:
    if len(s) <= max_chars:
        return s
    return s[: max_chars - 20] + "\n...<truncated>...\n"


@dataclass
class ToolResult:
    ok: bool
    payload: dict[str, Any]

    def to_tool_content(self) -> str:
        return json.dumps(self.payload, ensure_ascii=False)


class Tree(BaseModel):
    root: str = Field(..., description="Root folder to inspect.")
    max_depth: int = Field(3, ge=1, le=10, description="Max recursion depth.")
    max_entries_per_dir: int = Field(60, ge=10, le=500, description="Cap entries per directory.")


class ListFiles(BaseModel):
    root: str = Field(..., description="Root folder with manuals.")
    extensions: list[str] = Field(
        default_factory=lambda: [".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md"],
        description="Extensions to include.",
    )
    max_files: int = Field(300, ge=1, le=5000)


class Search(BaseModel):
    root: str = Field(..., description="Root folder with manuals.")
    query: str = Field(..., min_length=2, description="Search query (usually a fixed phrase).")
    mode: Literal["fixed", "regex"] = Field("fixed", description="Search mode.")
    extensions: list[str] = Field(
        default_factory=lambda: [
            "pdf",
            "docx",
            "pptx",
            "xlsx",
            "xlsm",
            "xltx",
            "xltm",
            "txt",
            "md",
        ],
        description="Which file extensions to include (with or without dot).",
    )
    max_results: int = Field(30, ge=1, le=200)
    context_chars: int = Field(180, ge=0, le=800)


class ReadTextFile(BaseModel):
    path: str = Field(..., description="Path to a text-like file within docs root.")
    max_chars: int = Field(8000, ge=100, le=200000)


class ReadPdfPages(BaseModel):
    path: str = Field(..., description="Path to a PDF within docs root.")
    page_start: int = Field(..., ge=1, description="1-based page start.")
    page_end: int = Field(..., ge=1, description="1-based page end (inclusive).")
    max_chars: int = Field(16000, ge=500, le=200000)


def _walk_tree_lines(root: str, max_depth: int, max_entries_per_dir: int) -> list[str]:
    root = _safe_abspath(root)
    lines: list[str] = [root]

    def rec(cur: str, depth: int):
        if depth > max_depth:
            return
        try:
            entries = sorted(os.listdir(cur))
        except Exception as e:
            lines.append("  " * depth + f"[error listing: {e}]")
            return

        shown = 0
        for name in entries:
            if name in {".git", ".venv", "venv", "__pycache__"}:
                continue
            path = os.path.join(cur, name)
            prefix = "  " * depth
            if os.path.isdir(path):
                lines.append(prefix + f"{name}/")
                shown += 1
                if shown >= max_entries_per_dir:
                    lines.append(prefix + "...<more entries truncated>...")
                    return
                rec(path, depth + 1)
            else:
                lines.append(prefix + name)
                shown += 1
                if shown >= max_entries_per_dir:
                    lines.append(prefix + "...<more entries truncated>...")
                    return

    rec(root, 1)
    return lines


def _list_files(root: str, extensions: Sequence[str], max_files: int) -> list[str]:
    root = _safe_abspath(root)
    exts = {e.lower() if e.startswith(".") else "." + e.lower() for e in extensions}
    files: list[str] = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in {".git", ".venv", "venv", "__pycache__"}]
        for fn in filenames:
            if len(files) >= max_files:
                return files
            ext = os.path.splitext(fn)[1].lower()
            if ext in exts:
                files.append(os.path.join(dirpath, fn))
    return files


def _have(cmd: str) -> bool:
    return shutil.which(cmd) is not None


def _find_ugrep_exe() -> str | None:
    override = (os.environ.get("AVIA_UGREP_PATH") or "").strip()
    if override and os.path.exists(override):
        return override

    local = os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "tools", "ugrep", "ugrep.exe")
    )
    if os.path.exists(local):
        return local

    return shutil.which("ugrep")


def _python_for_filters() -> str:
    override = (os.environ.get("AVIA_PYTHON") or "").strip()
    if override:
        return override
    return sys.executable


def _normalize_exts(exts: Sequence[str]) -> list[str]:
    out: list[str] = []
    seen = set()
    for e in exts or []:
        e = (e or "").strip().lower()
        if not e:
            continue
        if e.startswith("."):
            e = e[1:]
        if e not in seen:
            seen.add(e)
            out.append(e)
    return out


def _run_ugrep(
    root: str, query: str, mode: str, max_results: int, exts: Sequence[str]
) -> tuple[bool, str, str]:
    ugrep = _find_ugrep_exe()
    if not ugrep:
        return (
            False,
            "",
            "ugrep not found (set AVIA_UGREP_PATH or run scripts/bootstrap.ps1)",
        )

    project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    py = _python_for_filters()
    try:
        py_rel = os.path.relpath(py, project_root)
    except Exception:
        py_rel = py

    common = [
        ugrep,
        "--color=never",
        "-r",
        "-n",
        "-H",
        "--no-messages",
        "-m",
        str(max_results),
    ]
    mode_flags: list[str] = ["-F"] if mode == "fixed" else []

    wanted = set(_normalize_exts(exts))
    if not wanted:
        wanted = {"pdf", "docx", "pptx", "xlsx", "xlsm", "xltx", "xltm", "txt", "md"}

    passes: list[list[str]] = []
    if "pdf" in wanted:
        passes.append(
            common
            + [
                f"--filter=pdf:{py_rel} -m avia_agent.filter_stdin pdf",
                "-Opdf",
                *mode_flags,
                query,
                root,
            ]
        )
    if "docx" in wanted:
        passes.append(
            common
            + [
                f"--filter=docx:{py_rel} -m avia_agent.filter_stdin docx",
                "-Odocx",
                *mode_flags,
                query,
                root,
            ]
        )
    if "pptx" in wanted:
        passes.append(
            common
            + [
                f"--filter=pptx:{py_rel} -m avia_agent.filter_stdin pptx",
                "-Opptx",
                *mode_flags,
                query,
                root,
            ]
        )
    if wanted.intersection({"xlsx", "xlsm", "xltx", "xltm"}):
        passes.append(
            common
            + [
                f"--filter=xlsx,xlsm,xltx,xltm:{py_rel} -m avia_agent.filter_stdin xlsx",
                "-Oxlsx,xlsm,xltx,xltm",
                *mode_flags,
                query,
                root,
            ]
        )
    if wanted.intersection({"txt", "md"}):
        passes.append(common + ["-Otxt,md", *mode_flags, query, root])
    try:
        env = os.environ.copy()
        env["PYTHONPATH"] = project_root + (
            os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else ""
        )

        outs: list[str] = []
        errs: list[str] = []
        ok = True
        for cmd in passes:
            p = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                check=False,
                env=env,
                cwd=project_root,
            )
            if p.returncode not in (0, 1):
                ok = False
            if p.stdout:
                outs.append(p.stdout.strip())
            if p.stderr:
                errs.append(p.stderr.strip())

        out = "\n".join(o for o in outs if o)
        err = "\n".join(e for e in errs if e)
        return ok, out, err
    except Exception as e:
        return False, "", str(e)


def _parse_ugrep_lines(out: str, max_results: int) -> list[dict[str, str]]:
    results: list[dict[str, str]] = []
    for line in (out or "").splitlines():
        if len(results) >= max_results:
            break
        line = line.strip("\r\n")
        if not line:
            continue

        m = re.match(r"^(?P<path>.*):(?P<line>\d+):(?P<rest>.*)$", line)
        if not m:
            continue
        path = m.group("path").strip()
        line_no = m.group("line").strip()
        rest = m.group("rest")

        col = ""
        m2 = re.match(r"^(?P<col>\d+):(?P<text>.*)$", rest)
        if m2:
            col = m2.group("col").strip()
            text = m2.group("text").strip()
        else:
            text = rest.strip()

        loc = f"line {line_no}" + (f", col {col}" if col else "")
        results.append({"path": path, "location": loc, "snippet": text})
    return results


def _python_extract_pdf_pages(path: str, page_start: int, page_end: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    n = len(reader.pages)
    page_start = max(1, min(page_start, n))
    page_end = max(1, min(page_end, n))
    if page_end < page_start:
        page_start, page_end = page_end, page_start

    parts: list[str] = []
    for i in range(page_start - 1, page_end):
        page = reader.pages[i]
        txt = page.extract_text() or ""
        parts.append(f"\n--- page {i + 1} ---\n{txt}")
    return "\n".join(parts)


def _python_extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == ".pdf":
        return _python_extract_pdf_pages(path, 1, 999999)

    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text)

    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(path)
        chunks: list[str] = []
        for si, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text:
                    chunks.append(f"[slide {si}] {text}")
        return "\n".join(chunks)

    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        chunks: list[str] = []
        for sheet in wb.worksheets:
            chunks.append(f"[sheet {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None and str(v).strip() != ""]
                if vals:
                    chunks.append("\t".join(vals))
        return "\n".join(chunks)

    with open(path, "rb") as f:
        data = f.read(2000)
    raise ValueError(f"Unsupported file type for extraction: {ext}. First bytes: {data[:64]!r}")


def _python_search(
    root: str,
    query: str,
    mode: str,
    max_results: int,
    context_chars: int,
    extensions: Sequence[str],
) -> list[dict[str, str]]:
    root = _safe_abspath(root)
    if mode == "regex":
        rx = re.compile(query, flags=re.IGNORECASE)
    else:
        q = query.lower()

    exts = []
    for e in _normalize_exts(extensions):
        exts.append("." + e)
    if not exts:
        exts = [".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md"]

    candidates = _list_files(root, exts, max_files=100000)
    results: list[dict[str, str]] = []

    for path in candidates:
        if len(results) >= max_results:
            break

        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".pdf":
                from pypdf import PdfReader

                reader = PdfReader(path)
                for pi, page in enumerate(reader.pages, start=1):
                    if len(results) >= max_results:
                        break
                    txt = page.extract_text() or ""
                    hay = txt if mode == "regex" else txt.lower()
                    if mode == "regex":
                        m = rx.search(hay)
                        if not m:
                            continue
                        start, end = m.span()
                    else:
                        idx = hay.find(q)
                        if idx < 0:
                            continue
                        start, end = idx, idx + len(q)
                    left = max(0, start - context_chars)
                    right = min(len(txt), end + context_chars)
                    snippet = txt[left:right].replace("\r", " ").replace("\n", " ").strip()
                    results.append(
                        {
                            "path": path,
                            "location": f"page {pi}",
                            "snippet": snippet,
                        }
                    )
            else:
                txt = _python_extract_text(path)
                hay = txt if mode == "regex" else txt.lower()
                if mode == "regex":
                    m = rx.search(hay)
                    if not m:
                        continue
                    start, end = m.span()
                else:
                    idx = hay.find(q)
                    if idx < 0:
                        continue
                    start, end = idx, idx + len(q)
                left = max(0, start - context_chars)
                right = min(len(txt), end + context_chars)
                snippet = txt[left:right].replace("\r", " ").replace("\n", " ").strip()
                results.append({"path": path, "location": "", "snippet": snippet})
        except Exception:
            continue

    return results


def dispatch_tool(tool: Any, state: dict[str, Any]) -> ToolResult:
    docs_root = state.get("docs_root") or state.get("root") or ""
    docs_root = _safe_abspath(docs_root) if docs_root else ""

    if isinstance(tool, Tree):
        root = _safe_abspath(tool.root)
        if docs_root and not _is_within(docs_root, root) and root != docs_root:
            return ToolResult(
                False,
                {
                    "error": "root must be within docs_root",
                    "docs_root": docs_root,
                    "root": root,
                },
            )
        lines = _walk_tree_lines(root, tool.max_depth, tool.max_entries_per_dir)
        out = "\n".join(lines)
        return ToolResult(True, {"root": root, "tree": _truncate(out, 20000)})

    if isinstance(tool, ListFiles):
        root = _safe_abspath(tool.root)
        if docs_root and not _is_within(docs_root, root) and root != docs_root:
            return ToolResult(
                False,
                {
                    "error": "root must be within docs_root",
                    "docs_root": docs_root,
                    "root": root,
                },
            )
        files = _list_files(root, tool.extensions, tool.max_files)
        return ToolResult(
            True, {"root": root, "count": len(files), "files": files[: tool.max_files]}
        )

    if isinstance(tool, Search):
        root = _safe_abspath(tool.root)
        if docs_root and not _is_within(docs_root, root) and root != docs_root:
            return ToolResult(
                False,
                {
                    "error": "root must be within docs_root",
                    "docs_root": docs_root,
                    "root": root,
                },
            )

        backend = (os.environ.get("AVIA_SEARCH_BACKEND") or "auto").strip().lower()
        if backend == "auto":
            backend = "ugrep" if _find_ugrep_exe() else "python"

        if backend == "ugrep":
            ok, out, err = _run_ugrep(
                root, tool.query, tool.mode, tool.max_results, tool.extensions
            )
            parsed = _parse_ugrep_lines(out, tool.max_results) if ok else []
            payload = {
                "backend": "ugrep",
                "ugrep": _find_ugrep_exe(),
                "root": root,
                "query": tool.query,
                "ok": ok,
                "extensions": tool.extensions,
                "results": parsed,
                "raw": _truncate(out, 20000),
            }
            if err:
                payload["stderr"] = _truncate(err, 4000)
            return ToolResult(ok, payload)

        results = _python_search(
            root,
            tool.query,
            tool.mode,
            tool.max_results,
            tool.context_chars,
            tool.extensions,
        )
        return ToolResult(
            True,
            {
                "backend": "python",
                "root": root,
                "query": tool.query,
                "extensions": tool.extensions,
                "results": results,
            },
        )

    if isinstance(tool, ReadTextFile):
        path = _safe_abspath(tool.path)
        if docs_root and not _is_within(docs_root, path):
            return ToolResult(
                False,
                {
                    "error": "path must be within docs_root",
                    "docs_root": docs_root,
                    "path": path,
                },
            )
        try:
            with open(path, encoding="utf-8", errors="ignore") as f:
                txt = f.read()
            return ToolResult(True, {"path": path, "text": _truncate(txt, tool.max_chars)})
        except Exception as e:
            return ToolResult(False, {"path": path, "error": str(e)})

    if isinstance(tool, ReadPdfPages):
        path = _safe_abspath(tool.path)
        if docs_root and not _is_within(docs_root, path):
            return ToolResult(
                False,
                {
                    "error": "path must be within docs_root",
                    "docs_root": docs_root,
                    "path": path,
                },
            )
        try:
            txt = _python_extract_pdf_pages(path, tool.page_start, tool.page_end)
            return ToolResult(
                True,
                {
                    "path": path,
                    "page_start": tool.page_start,
                    "page_end": tool.page_end,
                    "text": _truncate(txt, tool.max_chars),
                },
            )
        except Exception as e:
            return ToolResult(False, {"path": path, "error": str(e)})

    return ToolResult(False, {"error": f"Unknown tool: {type(tool).__name__}"})
