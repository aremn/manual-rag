import os
import sys


def _read_text_file(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def _pdf_to_text(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        txt = page.extract_text() or ""
        parts.append(f"\n@@@PAGE {i}@@@\n{txt}")
    return "\n".join(parts)


def _docx_to_text(path: str) -> str:
    from docx import Document

    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _pptx_to_text(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    chunks: list[str] = []
    for slide_i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                chunks.append(f"[slide {slide_i}] {text}")
    return "\n".join(chunks)


def _xlsx_to_text(path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"\n--- sheet {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip() != ""]
            if vals:
                chunks.append("\t".join(vals))
    return "\n".join(chunks)


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}:
        return _read_text_file(path)
    if ext == ".pdf":
        return _pdf_to_text(path)
    if ext == ".docx":
        return _docx_to_text(path)
    if ext == ".pptx":
        return _pptx_to_text(path)
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _xlsx_to_text(path)
    raise SystemExit(f"Unsupported file type: {ext}")


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        print("usage: python -m avia_agent.extract_text <path>", file=sys.stderr)
        return 2
    path = argv[0]
    try:
        data = extract_text(path).encode("utf-8", errors="ignore")
        sys.stdout.buffer.write(data)
        return 0
    except Exception as e:
        print(f"error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
